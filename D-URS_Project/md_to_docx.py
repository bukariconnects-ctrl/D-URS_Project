"""
Convert academic_documentation.md (Arabic, RTL) to a formatted Word document.
Handles: headings, paragraphs, lists, tables, code blocks, horizontal rules, bold/italic, links.
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement


SRC = Path(__file__).parent / "academic_documentation.md"
DST = Path(__file__).parent / "academic_documentation.docx"

ARABIC_FONT = "Calibri"     # Latin font name; Arabic glyphs use complex-script font
CS_FONT = "Traditional Arabic"  # or "Arial" / "Calibri" - we'll use Arial-style for clarity
CS_FONT = "Arial"
MONO_FONT = "Consolas"


# --------------------------------------------------------------------------------
# Low-level helpers for RTL & font settings
# --------------------------------------------------------------------------------

def set_rtl(paragraph):
    """Mark a paragraph as right-to-left."""
    pPr = paragraph._p.get_or_add_pPr()
    bidi = pPr.find(qn('w:bidi'))
    if bidi is None:
        bidi = OxmlElement('w:bidi')
        pPr.append(bidi)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def set_run_font(run, size=12, bold=False, italic=False, color=None, mono=False):
    run.bold = bold
    run.italic = italic
    font_name = MONO_FONT if mono else ARABIC_FONT
    cs_font = MONO_FONT if mono else CS_FONT
    run.font.name = font_name
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    # Latin/Western font
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), cs_font)      # complex script (Arabic)
    # Mark run as RTL complex script
    rtl = rPr.find(qn('w:rtl'))
    if rtl is None:
        rtl = OxmlElement('w:rtl')
        rPr.append(rtl)
    # Complex-script size
    szCs = rPr.find(qn('w:szCs'))
    if szCs is None:
        szCs = OxmlElement('w:szCs')
        rPr.append(szCs)
    szCs.set(qn('w:val'), str(int(size * 2)))
    if bold:
        bCs = rPr.find(qn('w:bCs'))
        if bCs is None:
            bCs = OxmlElement('w:bCs')
            rPr.append(bCs)


# --------------------------------------------------------------------------------
# Inline markdown parsing (bold, italic, code, links)
# --------------------------------------------------------------------------------

INLINE_RE = re.compile(
    r'(\*\*[^*\n]+\*\*)'        # bold
    r'|(\*[^*\n]+\*)'           # italic
    r'|(`[^`\n]+`)'             # inline code
    r'|(\[[^\]]+\]\([^)]+\))'   # link
)


def add_inline(paragraph, text, base_size=12, base_bold=False):
    """Parse inline markdown and add styled runs to paragraph."""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos:m.start()])
            set_run_font(run, size=base_size, bold=base_bold)
        token = m.group(0)
        if token.startswith('**'):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, bold=True)
        elif token.startswith('*'):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=base_size, bold=base_bold, italic=True)
        elif token.startswith('`'):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=base_size - 1, bold=base_bold, mono=True,
                         color=RGBColor(0xC7, 0x25, 0x4E))
        elif token.startswith('['):
            # link [text](url) -> just keep text (styled as link color)
            lm = re.match(r'\[([^\]]+)\]\(([^)]+)\)', token)
            link_text = lm.group(1) if lm else token
            run = paragraph.add_run(link_text)
            set_run_font(run, size=base_size, bold=base_bold,
                         color=RGBColor(0x1F, 0x4E, 0x79))
            run.font.underline = True
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=base_size, bold=base_bold)


# --------------------------------------------------------------------------------
# Block-level handlers
# --------------------------------------------------------------------------------

def add_heading(doc, text, level):
    sizes = {1: 22, 2: 18, 3: 15, 4: 13}
    colors = {
        1: RGBColor(0x1F, 0x3A, 0x68),
        2: RGBColor(0x1F, 0x4E, 0x79),
        3: RGBColor(0x2E, 0x6F, 0x9E),
        4: RGBColor(0x44, 0x7A, 0xA8),
    }
    p = doc.add_paragraph()
    set_rtl(p)
    p.paragraph_format.space_before = Pt(12 if level > 1 else 18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    # remove inline markdown from heading (rare) and add plain styled run
    clean = re.sub(r'[*`]', '', text)
    run = p.add_run(clean)
    set_run_font(run, size=sizes.get(level, 12), bold=True, color=colors.get(level))


def add_paragraph_text(doc, text):
    p = doc.add_paragraph()
    set_rtl(p)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    add_inline(p, text, base_size=12)


def add_list_item(doc, text, ordered=False, indent_level=0, number=None):
    p = doc.add_paragraph()
    set_rtl(p)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.right_indent = Cm(0.5 + indent_level * 0.6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    bullet = f"{number}. " if ordered and number else "• "
    run = p.add_run(bullet)
    set_run_font(run, size=12, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    add_inline(p, text, base_size=12)


def add_code_block(doc, code_text):
    """Render fenced code block as a single shaded paragraph with monospace font."""
    p = doc.add_paragraph()
    # Code blocks are usually LTR (ASCII diagrams)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.right_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    # Shade background
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F4F4F4')
    pPr.append(shd)
    # Border
    pBdr = OxmlElement('w:pBdr')
    for side in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:space'), '4')
        b.set(qn('w:color'), 'D0D0D0')
        pBdr.append(b)
    pPr.append(pBdr)
    run = p.add_run(code_text)
    run.font.name = MONO_FONT
    run.font.size = Pt(10)
    rPr = run._element.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), MONO_FONT)
    rFonts.set(qn('w:hAnsi'), MONO_FONT)
    rFonts.set(qn('w:cs'), MONO_FONT)
    rPr.append(rFonts)


def add_table(doc, rows):
    """rows: list of list[str]; first row is header."""
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    # Set table to RTL
    tblPr = table._tbl.tblPr
    bidiVisual = OxmlElement('w:bidiVisual')
    tblPr.append(bidiVisual)

    for i, row_data in enumerate(rows):
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            text = row_data[j] if j < len(row_data) else ""
            cell.text = ""  # clear default paragraph
            p = cell.paragraphs[0]
            set_rtl(p)
            is_header = (i == 0)
            add_inline(p, text, base_size=11, base_bold=is_header)
            if is_header:
                # shade header
                tcPr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:fill'), '1F4E79')
                tcPr.append(shd)
                # white bold text already set; recolor runs to white
                for run in p.runs:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    run.bold = True


def add_horizontal_rule(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '8')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '1F4E79')
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_after = Pt(6)


# --------------------------------------------------------------------------------
# Main parser
# --------------------------------------------------------------------------------

def parse_markdown(md_text):
    """Yield ('block_type', data) tuples."""
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Fenced code block
        if stripped.startswith('```'):
            j = i + 1
            buf = []
            while j < len(lines) and not lines[j].strip().startswith('```'):
                buf.append(lines[j])
                j += 1
            yield ('code', '\n'.join(buf))
            i = j + 1
            continue

        # Horizontal rule
        if re.match(r'^-{3,}$', stripped) or re.match(r'^\*{3,}$', stripped):
            yield ('hr', None)
            i += 1
            continue

        # Heading
        h = re.match(r'^(#{1,6})\s+(.*)$', stripped)
        if h:
            level = len(h.group(1))
            yield ('heading', (level, h.group(2).strip()))
            i += 1
            continue

        # Table (line containing | and next line is separator)
        if '|' in line and i + 1 < len(lines) and re.match(r'^\s*\|?[\s:|-]+\|[\s:|-]+', lines[i+1]):
            rows = []
            # parse header row
            def split_row(row_line):
                parts = row_line.strip().strip('|').split('|')
                return [p.strip() for p in parts]
            rows.append(split_row(line))
            i += 2  # skip separator
            while i < len(lines) and '|' in lines[i] and lines[i].strip():
                rows.append(split_row(lines[i]))
                i += 1
            yield ('table', rows)
            continue

        # List item (unordered)
        ul = re.match(r'^(\s*)[-*+]\s+(.*)$', line)
        if ul:
            indent = len(ul.group(1)) // 2
            yield ('ul', (indent, ul.group(2).strip()))
            i += 1
            continue

        # List item (ordered)
        ol = re.match(r'^(\s*)(\d+)\.\s+(.*)$', line)
        if ol:
            indent = len(ol.group(1)) // 2
            yield ('ol', (indent, int(ol.group(2)), ol.group(3).strip()))
            i += 1
            continue

        # Blank line
        if not stripped:
            yield ('blank', None)
            i += 1
            continue

        # Paragraph (may span multiple lines until blank or block boundary)
        para_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt:
                break
            if re.match(r'^#{1,6}\s', nxt) or nxt.startswith('```') \
               or re.match(r'^-{3,}$', nxt) or re.match(r'^[-*+]\s', nxt) \
               or re.match(r'^\d+\.\s', nxt) or '|' in nxt:
                break
            para_lines.append(nxt)
            i += 1
        yield ('p', ' '.join(para_lines))


# --------------------------------------------------------------------------------
# Page setup
# --------------------------------------------------------------------------------

def setup_document(doc):
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        # Make section RTL
        sectPr = section._sectPr
        bidi = OxmlElement('w:bidi')
        sectPr.append(bidi)

    # Default style font
    style = doc.styles['Normal']
    style.font.name = ARABIC_FONT
    style.font.size = Pt(12)
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.append(rFonts)
    rFonts.set(qn('w:cs'), CS_FONT)
    rFonts.set(qn('w:ascii'), ARABIC_FONT)
    rFonts.set(qn('w:hAnsi'), ARABIC_FONT)


# --------------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------------

def main():
    md_text = SRC.read_text(encoding='utf-8')
    doc = Document()
    setup_document(doc)

    for kind, data in parse_markdown(md_text):
        if kind == 'heading':
            level, text = data
            add_heading(doc, text, level)
        elif kind == 'p':
            add_paragraph_text(doc, data)
        elif kind == 'code':
            add_code_block(doc, data)
        elif kind == 'hr':
            add_horizontal_rule(doc)
        elif kind == 'ul':
            indent, text = data
            add_list_item(doc, text, ordered=False, indent_level=indent)
        elif kind == 'ol':
            indent, number, text = data
            add_list_item(doc, text, ordered=True, indent_level=indent, number=number)
        elif kind == 'table':
            add_table(doc, data)
        elif kind == 'blank':
            pass

    doc.save(DST)
    print(f"Saved: {DST}")


if __name__ == '__main__':
    main()
